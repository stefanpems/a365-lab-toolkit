"""Renders the fictional knowledge of a demo-pack locale: .docx (python-docx), .pdf (fpdf2) and the personal .xlsx
(openpyxl). Input: a JSON file { "org": {...}, "folders": {...}, "documents": {...}, "personal": {...},
"packDocuments": [ {key, folder, format} ] } written by New-DemoKnowledge.ps1. Output: <out>/<folder>/<file> and
<out>/_personal/<file>, plus <out>/manifest.json. Block model: see the locale knowledge.json _comment.
Usage: python build_knowledge.py --input <json> --out <dir>
"""
import argparse
import json
import os
import re
from pathlib import Path

BOLD = re.compile(r"(\*\*[^*]+\*\*)")


def segments(text):
    for s in BOLD.split(str(text)):
        if s:
            yield (s[2:-2], True) if s.startswith("**") and s.endswith("**") else (s, False)


# --------------------------------------------------------------------------------------------------- DOCX
def build_docx(doc, org, path, slots=None, missing=None):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    slots = slots or {}
    missing = missing if missing is not None else set()
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2))
    width = sec.page_width - sec.left_margin - sec.right_margin
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Arial", Pt(11)
    banner = org.get("fictionalBanner", "")
    owner = f"{org.get('name', '')} – {org.get('department', '')}"

    hp = sec.header.paragraphs[0]
    hp.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
    r = hp.add_run(owner); r.font.size = Pt(8); r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    r = hp.add_run("\t" + banner); r.font.size = Pt(8); r.bold = True; r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

    fp = sec.footer.paragraphs[0]
    fp.alignment = 1
    def field(par, code):
        f = OxmlElement("w:fldSimple"); f.set(qn("w:instr"), code)
        rr = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "1"; rr.append(t); f.append(rr); par._p.append(f)
    r = fp.add_run(f"{banner} · "); r.font.size = Pt(8)
    field(fp, "PAGE"); r = fp.add_run(" / "); r.font.size = Pt(8); field(fp, "NUMPAGES")

    p = d.add_paragraph(); r = p.add_run(doc["title"]); r.bold = True; r.font.size = Pt(17); r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    if doc.get("subtitle"):
        p = d.add_paragraph(); r = p.add_run(doc["subtitle"]); r.italic = True; r.font.size = Pt(12)
    p = d.add_paragraph(); r = p.add_run(doc.get("meta") or owner); r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    def rich(par, text, size=None, italic=False):
        for s, b in segments(text):
            rr = par.add_run(s); rr.bold = b; rr.italic = italic
            if size: rr.font.size = Pt(size)

    for b in doc["blocks"]:
        kind = b[0]
        if kind in ("h1", "h2"):
            h = d.add_heading(level=1 if kind == "h1" else 2); rich(h, b[1])
        elif kind in ("p", "pi"):
            rich(d.add_paragraph(), b[1], italic=(kind == "pi"))
        elif kind == "slot":
            text = slots.get(b[1], "")
            if text:  # operator-supplied text, rendered invisible (white, 1 pt); never printed
                rr = d.add_paragraph().add_run(text); rr.font.size = Pt(1); rr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            else:
                missing.add(b[1])
        elif kind in ("bullets", "numbered"):
            for item in b[1]:
                rich(d.add_paragraph(style="List Bullet" if kind == "bullets" else "List Number"), item)
        elif kind == "table":
            headers, rows, weights = b[1], b[2], b[3]
            t = d.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
            total = float(sum(weights)) or 1.0
            for i, h in enumerate(headers):
                c = t.rows[0].cells[i]; c.text = ""; rich(c.paragraphs[0], f"**{h}**", size=9)
            for row in rows:
                cells = t.add_row().cells
                for i, v in enumerate(row):
                    cells[i].text = ""; rich(cells[i].paragraphs[0], v, size=9)
            for i, w in enumerate(weights):
                for row in t.rows:
                    row.cells[i].width = int(width * (w / total))
            d.add_paragraph()
        else:
            raise ValueError(f"unknown block {kind}")
    d.core_properties.title = doc["title"]
    d.core_properties.author = owner
    d.core_properties.comments = banner
    d.save(path)


# --------------------------------------------------------------------------------------------------- PDF
def pdf_font(pdf):
    for cand in (r"C:\Windows\Fonts\arial.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/Library/Fonts/Arial.ttf"):
        if os.path.exists(cand):
            bold = cand.replace("arial.ttf", "arialbd.ttf").replace("DejaVuSans.ttf", "DejaVuSans-Bold.ttf")
            pdf.add_font("Body", "", cand)
            pdf.add_font("Body", "B", bold if os.path.exists(bold) else cand)
            pdf.add_font("Body", "I", cand)
            return "Body", False
    return "Helvetica", True  # core font: latin-1 only


def build_pdf(doc, org, path):
    from fpdf import FPDF

    banner = org.get("fictionalBanner", "")
    owner = f"{org.get('name', '')} – {org.get('department', '')}"

    class P(FPDF):
        def header(self):
            self.set_font(font, "B", 8); self.set_text_color(0xC0, 0, 0)
            self.cell(0, 5, clean(f"{owner}   ·   {banner}"), align="R", new_x="LMARGIN", new_y="NEXT"); self.ln(2)

        def footer(self):
            self.set_y(-12); self.set_font(font, "", 8); self.set_text_color(0x7F, 0x7F, 0x7F)
            self.cell(0, 5, clean(f"{banner} · {self.page_no()}/{{nb}}"), align="C")

    pdf = P(format="A4"); pdf.alias_nb_pages()
    font, core = pdf_font(pdf)
    def clean(s):
        s = str(s)
        if core:
            s = s.replace("–", "-").replace("—", "-").replace("€", "EUR").replace("“", '"').replace("”", '"').replace("’", "'").replace("…", "...")
            s = s.encode("latin-1", "replace").decode("latin-1")
        return s
    def md(s):  # fpdf2 markdown bold uses the same ** syntax
        return clean(s)
    pdf.set_margins(18, 15, 18); pdf.set_auto_page_break(True, 18); pdf.add_page()
    pdf.set_font(font, "B", 16); pdf.set_text_color(0x1F, 0x38, 0x64); pdf.multi_cell(0, 8, clean(doc["title"]), new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(0, 0, 0)
    if doc.get("subtitle"):
        pdf.set_font(font, "I", 11); pdf.multi_cell(0, 6, clean(doc["subtitle"]), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font(font, "", 9); pdf.set_text_color(0x59, 0x59, 0x59); pdf.multi_cell(0, 5, clean(doc.get("meta") or owner), new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(0, 0, 0); pdf.ln(3)
    for b in doc["blocks"]:
        kind = b[0]
        if kind in ("h1", "h2"):
            pdf.ln(2); pdf.set_font(font, "B", 13 if kind == "h1" else 11); pdf.multi_cell(0, 6, clean(b[1]), new_x="LMARGIN", new_y="NEXT")
        elif kind in ("p", "pi"):
            pdf.set_font(font, "I" if kind == "pi" else "", 10); pdf.multi_cell(0, 5, md(b[1]), markdown=True, new_x="LMARGIN", new_y="NEXT"); pdf.ln(1)
        elif kind == "slot":
            continue  # operator slots are only supported in .docx documents
        elif kind in ("bullets", "numbered"):
            pdf.set_font(font, "", 10)
            for i, item in enumerate(b[1], 1):
                pdf.multi_cell(0, 5, md(("• " if kind == "bullets" and not core else ("- " if kind == "bullets" else f"{i}. ")) + item), markdown=True, new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1)
        elif kind == "table":
            headers, rows, weights = b[1], b[2], b[3]
            total = float(sum(weights)) or 1.0
            widths = [pdf.epw * w / total for w in weights]
            pdf.set_font(font, "", 8)
            with pdf.table(col_widths=widths, text_align="LEFT", line_height=4.5, first_row_as_headings=True) as t:
                r = t.row()
                for h in headers:
                    r.cell(clean(h))
                for row in rows:
                    r = t.row()
                    for v in row:
                        r.cell(md(v))
            pdf.ln(2)
    pdf.set_title(clean(doc["title"])); pdf.set_author(clean(owner)); pdf.set_subject(clean(banner))
    pdf.output(str(path))


# --------------------------------------------------------------------------------------------------- XLSX
def build_xlsx(spec, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook(); ws = wb.active; ws.title = spec.get("sheet", "Sheet1")[:31]
    ws["A1"] = spec.get("banner", ""); ws["A1"].font = Font(name="Arial", size=11, bold=True, color="C00000")
    ws.append([]); ws.append(spec["headers"])
    for c in ws[3]:
        c.font = Font(name="Arial", size=10, bold=True); c.fill = PatternFill("solid", start_color="D9E2F3")
    for row in spec["rows"]:
        ws.append(row)
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row):
        for c in row:
            c.font = Font(name="Arial", size=10); c.alignment = Alignment(vertical="top", wrap_text=True)
    for i, w in enumerate((28, 18, 26, 34, 18, 44)[: len(spec["headers"])]):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A4"
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    data = json.loads(Path(a.input).read_text(encoding="utf-8"))
    org, out = data["org"], Path(a.out)
    slots = data.get("slots") or {}
    manifest = []
    for pd in data["packDocuments"]:
        doc = data["documents"][pd["key"]]
        folder = out / data["folders"][pd["folder"]]
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / doc["file"]
        missing = set()
        if pd["format"] == "pdf":
            build_pdf(doc, org, path)
        else:
            build_docx(doc, org, path, slots, missing)
        manifest.append({"key": pd["key"], "folder": data["folders"][pd["folder"]], "file": doc["file"], "format": pd["format"],
                         "sensitivityLabel": pd.get("sensitivityLabel"), "poisoned": bool(pd.get("poisoned")),
                         "slotsMissing": sorted(missing)})
        print(f"written: {path.relative_to(out)}" + (f"  (operator slot not filled: {', '.join(sorted(missing))})" if missing else ""))
    for key, spec in (data.get("personal") or {}).items():
        folder = out / "_personal"; folder.mkdir(parents=True, exist_ok=True)
        build_xlsx(spec, folder / spec["file"])
        manifest.append({"key": key, "folder": "_personal", "file": spec["file"], "format": "xlsx", "sensitivityLabel": None, "poisoned": False, "slotsMissing": []})
        print(f"written: _personal/{spec['file']}")
    (out / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"documents: {len(manifest)}")


if __name__ == "__main__":
    main()
